from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # --- THEME COLORS (Azure Style) ---
    AZURE_BLUE = RGBColor(0, 120, 212)
    AZURE_LIGHT = RGBColor(222, 236, 249)
    DARK_GREY = RGBColor(80, 80, 80)
    
    def set_title_format(slide, text):
        title = slide.shapes.title
        title.text = text
        title.text_frame.paragraphs[0].font.color.rgb = AZURE_BLUE
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.name = "Segoe UI"

    def add_bullet(tf, text, level=0):
        p = tf.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(18)
        p.font.name = "Segoe UI"

    # --- SLIDE 1: TITLE ---
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Hybrid DevOps Architecture & Observability"
    title.text_frame.paragraphs[0].font.color.rgb = AZURE_BLUE
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle.text = "Target Platform: Hybrid (On-Prem + Azure)\nTools: Azure DevOps, Sentinel, Monitor, Terraform, Ansible"

    # --- SLIDE 2: EXECUTIVE SUMMARY ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    set_title_format(slide, "Executive Summary")
    
    tf = slide.placeholders[1].text_frame
    add_bullet(tf, "Objective: Unified CI/CD & Observability for Hybrid Infrastructure.")
    add_bullet(tf, "Primary Scope:", 0)
    add_bullet(tf, "Environments: Dev, Test, Stage, PreProd, Prod", 1)
    add_bullet(tf, "Hosting: Azure Resources + On-Premise Servers", 1)
    add_bullet(tf, "Key Pillars:", 0)
    add_bullet(tf, "IaC: Terraform (Infra) + Ansible (Config)", 1)
    add_bullet(tf, "Security: Defender for Cloud + Sentinel (SIEM/SOAR)", 1)
    add_bullet(tf, "Observability: Full stack monitoring (App Insights to Log Analytics)", 1)

    # --- SLIDE 3: ARCHITECTURE OVERVIEW ---
    slide_layout = prs.slide_layouts[1] # Using Title & Content
    slide = prs.slides.add_slide(slide_layout)
    set_title_format(slide, "High-Level Logical Architecture")
    
    # Add a text box describing the flow since we can't embed the image dynamically without a file
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(9)
    height = Inches(5)
    
    textbox = slide.shapes.add_textbox(left, top, width, height)
    tf = textbox.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = "[PLACEHOLDER FOR ARCHITECTURE DIAGRAM]\n"
    p.font.bold = True
    p.font.color.rgb = AZURE_BLUE
    
    p = tf.add_paragraph()
    p.text = "Logical Flow:"
    p.font.bold = True
    
    flows = [
        "1. PLAN: Azure Boards tracks work items.",
        "2. CODE: Azure Repos triggers Pipelines.",
        "3. SCAN: SonarQube checks quality; Key Vault provides secrets.",
        "4. DEPLOY: Terraform provisions Infra; Ansible configures OS.",
        "5. MONITOR: Agents send logs to Azure Monitor/Sentinel.",
        "6. SECURE: Defender protects workloads; Sentinel correlates threats."
    ]
    for item in flows:
        p = tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(16)

    # --- SLIDE 4: CI/CD TOOLCHAIN ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title_format(slide, "CI/CD & Automation Toolchain")
    tf = slide.placeholders[1].text_frame
    
    add_bullet(tf, "Orchestration: Azure Pipelines (YAML Multi-stage)")
    add_bullet(tf, "Infrastructure as Code (IaC):")
    add_bullet(tf, "Terraform: State managed in Azure Storage. Provisions VNETs, VMs, AKS.", 1)
    add_bullet(tf, "Ansible: Configures Middleware, Patches, App Deployment on VMs.", 1)
    add_bullet(tf, "Python: Custom automation scripts.", 1)
    add_bullet(tf, "Governance:")
    add_bullet(tf, "ServiceNow: Change Management gating.", 1)
    add_bullet(tf, "Azure Key Vault: Secret injection at runtime.", 1)
    add_bullet(tf, "Azure Test Plans: Manual & Automated testing.", 1)

    # --- SLIDE 5: OBSERVABILITY & AIOPS ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title_format(slide, "Observability Strategy (The Eyes)")
    tf = slide.placeholders[1].text_frame
    
    add_bullet(tf, "Application Layer:")
    add_bullet(tf, "Azure Application Insights: APM, Distributed Tracing.", 1)
    add_bullet(tf, "Azure Workbooks: Custom analytics visualization.", 1)
    add_bullet(tf, "Infrastructure Layer:")
    add_bullet(tf, "Azure Monitor Ops Insight: VM & Host health.", 1)
    add_bullet(tf, "Azure Container Log Analytics: Kubernetes/Container metrics.", 1)
    add_bullet(tf, "Service Map: Dependency modeling (Server-to-Server connections).", 1)
    add_bullet(tf, "AIOps:")
    add_bullet(tf, "Azure Monitor Insight: Intelligent anomaly detection.", 1)

    # --- SLIDE 6: SECURITY ARCHITECTURE ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title_format(slide, "Security & Sentinel (The Brain)")
    tf = slide.placeholders[1].text_frame
    
    add_bullet(tf, "SIEM & SOAR:")
    add_bullet(tf, "Azure Sentinel: Ingests logs from all Hybrid sources.", 1)
    add_bullet(tf, "Sentinel AIOps: ML-driven threat detection & correlation.", 1)
    add_bullet(tf, "Cloud Security Posture Management (CSPM):")
    add_bullet(tf, "Microsoft Defender for Cloud: Protects Azure & On-Prem workloads.", 1)
    add_bullet(tf, "Identity:")
    add_bullet(tf, "Azure Key Vault: Centralized secret management.", 1)

    # --- SLIDE 7: REPOSITORY STRUCTURE ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title_format(slide, "Roadmap & Repo Structure")
    tf = slide.placeholders[1].text_frame
    
    add_bullet(tf, "Repository: kkarthee/azure_architect")
    add_bullet(tf, "Folder Structure:")
    add_bullet(tf, "/iac - Terraform modules & Ansible playbooks", 1)
    add_bullet(tf, "/pipelines - YAML Pipeline definitions", 1)
    add_bullet(tf, "/src - Application source code", 1)
    add_bullet(tf, "/docs - Architecture diagrams & decision records", 1)
    add_bullet(tf, "Next Steps:", 0)
    add_bullet(tf, "1. Initialize Terraform Remote State.", 1)
    add_bullet(tf, "2. Configure Azure DevOps Service Connections (Azure & On-Prem).", 1)
    add_bullet(tf, "3. Deploy Sentinel Data Connectors.", 1)

    prs.save('Hybrid_DevOps_Architecture.pptx')
    print("Presentation saved as Hybrid_DevOps_Architecture.pptx")

if __name__ == "__main__":
    create_presentation()